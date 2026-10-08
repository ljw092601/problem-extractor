# -*- coding: utf-8 -*-
"""
problems_index.json 을 이용해, 원하는 문제 번호만 골라
학원 양식(수.hwp, template.py)에 맞춘 2단 Word(.docx) 문서로 조립한다.

입력은 '페이지:번호' 형식 (헷갈림 방지):
    python build_docx.py 2:5 2:7 3:1-4        # 2페이지 5·7번, 3페이지 1~4번
    python build_docx.py 1 5 9 16 23          # 번호가 유일하면 번호만도 가능
    python build_docx.py 2:5 3:1-4 -o 오답노트.docx   # 출력 파일명 지정
    python build_docx.py 1 5 9 --cols 1       # 1단으로

각 문제는 원본에서 고해상도로 잘라 PNG로 만든 뒤, 단 폭에 맞춰 이미지로 삽입한다.
Word가 좌단을 채우고 우단으로 자동으로 넘긴다.
프로그램(선생님용 앱)에서는 build(items, out_path) 를 직접 부른다.
"""
import io, os
from docx import Document
from docx.shared import Mm, Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import selector
import template as T

HERE = os.path.dirname(os.path.abspath(__file__))


def set_columns(section, num, gap_mm):
    """섹션을 num단으로 설정하고 단 사이 구분선을 켠다."""
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(gap_mm / 25.4 * 1440)))   # mm -> twips
    cols.set(qn("w:equalWidth"), "1")
    cols.set(qn("w:sep"), "1" if num > 1 else "0")            # 단 사이 가운데 구분선


def add_border(paragraph, side, sz=24, space=2):
    """문단 위(top) 또는 아래(bottom)에 굵은 가로선. sz: 1/8pt 단위(24=3pt), space: 글자와 선 사이(pt)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), "000000")
    pBdr.append(b)
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is not None:
        pStyle.addnext(pBdr)             # pStyle 다음(스키마상 올바른 위치)
    else:
        pPr.insert(0, pBdr)


def add_page_number(paragraph):
    """문단에 PAGE 필드(현재 페이지 번호)를 넣는다."""
    run = paragraph.add_run()
    set_font(run)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)


def set_font(run, size=T.HF_FONT_PT, name=T.HF_FONT_NAME):
    """한글까지 적용되도록 eastAsia 글꼴도 함께 지정."""
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), name)


WORD_PIC_EXTRA_MM = 1.5           # Word 가 그림 줄에 더하는 여분 높이(실측)


def add_blank_line(doc):
    """간격용 빈 줄: 10pt, 줄 높이 고정 16pt (수.hwp 본문과 같은 한 줄). 엔터로 추가하면 같은 모양이 복사된다."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = Pt(T.LINE_PT)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    # 문단 기호의 글자 크기 (pPr 의 마지막 자식 rPr) → 이 줄에서 치는 글자도 10pt
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(T.BODY_FONT_PT * 2)); rPr.append(sz)
    p._p.get_or_add_pPr().append(rPr)
    return p


def zero_spacing(paragraph):
    pf = paragraph.paragraph_format
    pf.space_before = pf.space_after = Pt(0)


def setup_page(doc, cols, title):
    """학원 양식: A4, 여백, 단, 머리말(파일 이름 + 페이지 번호 + 굵은 선), 꼬리말(굵은 선 + 페이지 번호)."""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(T.PAGE_W_MM), Mm(T.PAGE_H_MM)
    sec.left_margin = sec.right_margin = Mm(T.MARGIN_X_MM)
    sec.top_margin = Mm(T.BODY_TOP_MM)
    sec.bottom_margin = Mm(T.PAGE_H_MM - T.BODY_BOT_MM)
    sec.header_distance = sec.footer_distance = Mm(T.MARGIN_Y_MM)
    set_columns(sec, cols, T.COL_GAP_MM)

    # 머리말: 왼쪽 파일 이름, 오른쪽 끝 페이지 번호, 아래 굵은 선
    hp = sec.header.paragraphs[0]
    zero_spacing(hp)
    tabs = hp.paragraph_format.tab_stops
    for pos in (Twips(4680), Twips(9360)):      # Word 기본 머리말 스타일의 가운데/오른쪽 탭 제거
        tabs.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
    tabs.add_tab_stop(Mm(T.BODY_W_MM), WD_TAB_ALIGNMENT.RIGHT)
    set_font(hp.add_run(title + "\t"))
    add_page_number(hp)
    add_border(hp, "bottom")

    # 꼬리말: 위 굵은 선, 가운데 페이지 번호
    fp = sec.footer.paragraphs[0]
    zero_spacing(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_border(fp, "top")
    add_page_number(fp)


def build(items, out_path, cols=T.COLS, per_col=T.PER_COL):
    """items: [{"pix": 잘라 둔 문제(fitz.Pixmap, RENDER_DPI), "label": "[ 2:5 ]"}] → out_path(.docx) 저장."""
    doc = Document()
    setup_page(doc, cols, os.path.basename(out_path))   # 머리말에 파일 이름 (수.hwp 와 동일)
    add_items(doc, items, cols, per_col)
    doc.save(out_path)


def append(items, path, cols=T.COLS, per_col=T.PER_COL):
    """이미 만든 .docx 끝에 새 쪽부터 문제를 이어 붙여 같은 파일에 저장.
    새 구역(새 쪽에서 시작)을 추가해 단 수를 맞춘다. 머리말·꼬리말·쪽 번호는 앞 구역을 그대로 잇는다.
    선생님이 Word 에서 고친 내용은 그대로 둔다."""
    from docx.enum.section import WD_SECTION
    doc = Document(path)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_columns(sec, cols, T.COL_GAP_MM)
    add_items(doc, items, cols, per_col)
    tmp = path + ".tmp-append"
    doc.save(tmp)
    os.replace(tmp, path)


def add_items(doc, items, cols, per_col):
    """문서 끝에 문제 그림(+칸 맞춤 빈 줄)을 차례로 넣는다."""
    slot_h_mm = T.slot_height_mm(per_col)

    for it in items:
        pix = it["pix"]
        img = io.BytesIO(pix.tobytes("png"))
        w_mm, h_mm = T.fit_size_mm(pix.width / T.RENDER_DPI * 25.4, pix.height / T.RENDER_DPI * 25.4,
                                   cols, per_col)

        if per_col:
            # 칸 배치: 문제 아래를 빈 줄(엔터)로 채워 칸 높이를 맞춘다.
            # 선생님이 빈 줄을 지우거나 엔터를 쳐서 간격을 직접 조정할 수 있다.
            # 빈 줄 수는 칸을 넘지 않게 내림 → 다음 문제는 자기 칸 맨 위(또는 다음 단)에서 시작
            pic_p = doc.add_paragraph()
            pf = pic_p.paragraph_format
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pic_p.add_run().add_picture(img, width=Mm(w_mm))
            k = T.slots_for(h_mm, per_col)                   # 긴 문제는 칸을 여러 개 차지
            for _ in range(T.blank_lines(k * slot_h_mm - h_mm - WORD_PIC_EXTRA_MM)):
                add_blank_line(doc)
            continue

        if T.SHOW_LABEL and it.get("label"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True    # 라벨만 단 끝에 남지 않게
            run = p.add_run(it["label"])
            run.bold = True
            set_font(run, size=T.LABEL_FONT_PT)
            run.font.color.rgb = RGBColor(*T.LABEL_COLOR)

        # 이어 붙이기: 문제 사이 10pt 간격으로 차례로 흘려 넣음
        pic_p = doc.add_paragraph()
        pic_p.paragraph_format.space_after = Pt(10)
        pic_p.add_run().add_picture(img, width=Mm(w_mm))


def main():
    args = selector.cli_args("selected_problems.docx")
    items = selector.cli_items(args.numbers)
    build(items, os.path.join(HERE, args.out), args.cols, args.per_col)
    selector.cli_report(args, items)


if __name__ == "__main__":
    main()
